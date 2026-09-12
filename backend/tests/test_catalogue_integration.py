"""INT-01/02: synthetic workbooks only; no embeddings, private sources, or network."""
import csv
from contextlib import redirect_stderr, redirect_stdout
import importlib.util
import io
from pathlib import Path
import shutil
import unittest
from unittest.mock import Mock, patch
import uuid
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import Workbook, load_workbook
from app.services import ingest
from test_parcours_ux import PARCOURS_ROOT, PageParser, synthetic_case


def cache_formula(path, coordinate, text):
    """Simulate an Excel-calculated cached string in a synthetic fixture."""
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with zipfile.ZipFile(io.BytesIO(path.read_bytes())) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    sheet = ET.fromstring(files["xl/worksheets/sheet1.xml"])
    cell = next(c for c in sheet.iter(ns + "c") if c.attrib["r"] == coordinate)
    cell.set("t", "str")
    cell.find(ns + "v").text = text
    files["xl/worksheets/sheet1.xml"] = ET.tostring(sheet)
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)


class SyntheticFiles(unittest.TestCase):
    def setUp(self):
        self.root = Path.cwd() / f".synthetic-int-{uuid.uuid4().hex}"
        self.root.mkdir()
        self.addCleanup(shutil.rmtree, self.root)

    def workbook(self, sheets, name="synthetic.xlsx"):
        path = self.root / name
        wb = Workbook()
        wb.remove(wb.active)
        try:
            for title, rows in sheets:
                sheet = wb.create_sheet(title)
                for row in rows:
                    sheet.append(row)
            wb.save(path)
        finally:
            wb.close()
        return path


class CatalogueImportTests(SyntheticFiles):
    loader = ingest

    def load(self, sheets):
        return self.loader._load_xlsx(str(self.workbook(sheets)))

    def test_sheet1_wins_over_leading_audit_and_proposal(self):
        headers = ["use_case_id", "rag_text_auto", "guardrails", "audit_note", "domaine_label_fr", "micro_theme"]
        docs = self.load([
            ("AUDIT", [["private-review-only"], ["not served"]]),
            ("BASE_PROPOSEE", [headers, ["PROPOSAL", "proposal text", "", "", "", ""]]),
            ("Sheet1", [headers, [" 0001 ", " Texte é\nexact & <ici> ", "Garde-fou.", "review", "Domaine", " Micro-thème é "]]),
        ])
        self.assertEqual(len(docs), 1)
        self.assertEqual(docs[0].page_content, "Texte é\nexact & <ici>")
        self.assertEqual(docs[0].metadata["use_case_id"], "0001")
        self.assertEqual(docs[0].metadata["guardrails"], "Garde-fou.")
        self.assertEqual(docs[0].metadata["domaine_label_fr"], "Domaine")
        self.assertEqual(docs[0].metadata["micro_theme"], "Micro-thème é")
        self.assertNotIn("audit_note", docs[0].metadata)
        self.assertEqual(docs[0].metadata["sheet"], "Sheet1")

    def test_unambiguous_legacy_catalogue_is_supported(self):
        docs = self.load([
            ("Audit", [["notes"]]),
            ("Catalogue", [["use_case_id", "rag_text_auto"], ["SYN-01", "Fictif"]]),
        ])
        self.assertEqual(docs[0].metadata["sheet"], "Catalogue")

    def test_ambiguous_or_proposal_only_catalogues_fail(self):
        rows = [["use_case_id", "rag_text_auto"], ["SYN-01", "Fictif"]]
        for sheets in ([("A", rows), ("B", rows)], [("BASE_PROPOSEE", rows)]):
            with self.subTest(sheets=len(sheets)), self.assertRaises(ValueError):
                self.load(sheets)

    def test_malformed_sheet1_does_not_fall_back(self):
        with self.assertRaisesRegex(ValueError, "headers missing"):
            self.load([
                ("Sheet1", [["private-unexpected-header"]]),
                ("Catalogue", [["use_case_id", "rag_text_auto"], ["SYN-01", "Fictif"]]),
            ])

    def test_required_and_duplicate_headers_fail_without_values_in_error(self):
        for headers in (["rag_text_auto"], ["use_case_id"], ["use_case_id", "rag_text_auto", "rag_text_auto"]):
            with self.subTest(headers=headers), self.assertRaises(ValueError) as error:
                self.load([("Sheet1", [headers, ["SENSITIVE-SYNTHETIC", "secret-synthetic", "x"]])])
            self.assertNotIn("SENSITIVE-SYNTHETIC", str(error.exception))
            self.assertNotIn("secret-synthetic", str(error.exception))

    def test_duplicate_blank_ids_and_partial_rows_fail(self):
        headers = ["use_case_id", "rag_text_auto", "guardrails"]
        for rows in (
            [["SYN-01", "text"], [" SYN-01 ", "other"]],
            [["Syn-01", "text"], ["SYN-01", "other"]],
            [[None, "text"]], [["SYN-01", None]], [[None, None, "guardrail"]],
            [["", "", ""]], [["SYN-01", "   "]],
        ):
            with self.subTest(rows=rows), self.assertRaises(ValueError):
                self.load([("Sheet1", [headers, *rows])])

    def test_mixed_case_id_is_preserved_without_case_collision(self):
        docs = self.load([("Sheet1", [["use_case_id", "rag_text_auto"], [" Syn-01 ", "Fictif"]])])
        self.assertEqual(docs[0].metadata["use_case_id"], "Syn-01")

    def test_truly_blank_rows_are_ignored(self):
        docs = self.load([("Sheet1", [
            ["use_case_id", "rag_text_auto"], [None, None], ["SYN-01", "text"], [None, None],
        ])])
        self.assertEqual(len(docs), 1)
        self.assertEqual(docs[0].metadata["row_index"], 3)

    def test_missing_formula_caches_in_any_served_column_fail(self):
        for column in (0, 1, 2):
            row = ["SYN-01", "Fictif", "Garde-fou"]
            row[column] = '="synthetic-cache"'
            with self.subTest(column=column), self.assertRaisesRegex(ValueError, "formula cache"):
                self.load([("Sheet1", [["use_case_id", "rag_text_auto", "guardrails"], row])])

    def test_cached_formula_text_is_served_verbatim(self):
        path = self.workbook([("Sheet1", [["use_case_id", "rag_text_auto"], ["SYN-01", '="é fictif"']])])
        cache_formula(path, "B2", "é fictif\nSuite intacte")
        docs = self.loader._load_xlsx(str(path))
        self.assertEqual(docs[0].page_content, "é fictif\nSuite intacte")

    def test_excel_errors_fail_in_served_columns_not_audit_columns(self):
        for column in (0, 1, 2):
            row = ["SYN-01", "Fictif", "Garde-fou"]
            row[column] = "#VALUE!"
            with self.subTest(column=column), self.assertRaisesRegex(ValueError, "Excel error"):
                self.load([("Sheet1", [["use_case_id", "rag_text_auto", "guardrails"], row])])
        docs = self.load([("Sheet1", [
            ["use_case_id", "rag_text_auto", "audit"], ["SYN-01", "Fictif", "#VALUE!"],
        ])])
        self.assertNotIn("audit", docs[0].metadata)

    def test_populated_unnamed_columns_fail(self):
        with self.assertRaisesRegex(ValueError, "Unnamed"):
            self.load([("Sheet1", [["use_case_id", "rag_text_auto", None], ["SYN-01", "Fictif", "x"]])])

    def test_workbooks_closed_on_validation_failure_and_second_open_failure(self):
        path = self.workbook([("Sheet1", [["use_case_id"], ["SYN-01"]])])
        books = []

        def opening(*args, **kwargs):
            wb = load_workbook(*args, **kwargs)
            wb.close = Mock(wraps=wb.close)
            books.append(wb)
            return wb

        with patch.object(self.loader, "load_workbook", side_effect=opening), self.assertRaises(ValueError):
            self.loader._load_xlsx(str(path))
        self.assertEqual(len(books), 2)
        for wb in books:
            wb.close.assert_called_once()
        wb = load_workbook(path, read_only=True)
        wb.close = Mock(wraps=wb.close)
        with patch.object(self.loader, "load_workbook", side_effect=[wb, OSError("synthetic")]), \
                self.assertRaises(OSError):
            self.loader._load_xlsx(str(path))
        wb.close.assert_called_once()

    def test_real_langchain_xlsx_documents_convert_to_real_haystack(self):
        from haystack import Document
        from langchain_core.documents import Document as LangChainDocument
        path = self.workbook([("Sheet1", [["use_case_id", "rag_text_auto"], ["SYN-01", "Fictif é"]])])
        docs = self.loader.load_and_split_documents(str(path))
        self.assertIsInstance(docs[0], LangChainDocument)
        converted = self.loader._lc_to_haystack_docs(docs, str(path))
        self.assertIsInstance(converted[0], Document)
        self.assertEqual(converted[0].content, "Fictif é")
        self.assertEqual(converted[0].meta["use_case_id"], "SYN-01")

    def test_real_validate_only_path_reads_synthetic_xlsx_without_index_calls(self):
        from app.scripts import index_documents
        path = self.workbook([("Sheet1", [["use_case_id", "rag_text_auto"], ["SYN-01", "Fictif é"]])])
        stdout = io.StringIO()
        with patch.object(index_documents, "get_document_store", side_effect=AssertionError("No Chroma")), \
                patch.object(index_documents, "clear_all_documents", side_effect=AssertionError("No clear")), \
                patch.object(index_documents, "index_documents_haystack", side_effect=AssertionError("No embeddings")), \
                redirect_stdout(stdout):
            index_documents.main(["--validate-only", str(path)])
        self.assertIn("1 document(s)", stdout.getvalue())
        self.assertNotIn("Fictif é", stdout.getvalue())
        self.assertNotIn("SYN-01", stdout.getvalue())


class SafeGeneratorTests(SyntheticFiles):
    @classmethod
    def setUpClass(cls):
        source = PARCOURS_ROOT / "pipeline" / "genere.py"
        if not source.is_file():
            raise unittest.SkipTest("Set PARCOURS_SOURCE_ROOT for synthetic generator integration")
        spec = importlib.util.spec_from_file_location("safe_synthetic_generator", source)
        cls.generator = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.generator)

    def setUp(self):
        super().setUp()
        self.enterContext(patch.dict(self.generator.os.environ))
        self.generator.os.environ.pop("AVOULIA_APP_URL", None)
        case = synthetic_case()
        self.case = case
        self.record = {k: v for k, v in case.items() if not isinstance(v, list)}
        self.record.update(
            questions_qualification="|".join(case["questions"]),
            prerequis_donnees="|".join(case["prereqs"]),
            declencheurs_typiques="Déclencheur fictif",
            rag_text_auto="Recherche entièrement fictive",
        )
        self.headers = list(self.record)
        self.rows = [self.headers, list(self.record.values())]
        self.source = self.workbook([("Sheet1", self.rows)])
        self.mapping = self.write_mapping([["SYNTHETIC-UX", "synthetic01", "/action/synthetic01/"]])
        self.output = self.root / "new-stage"

    def write_mapping(self, rows, headers=("use_case_id", "hash", "url")):
        path = self.root / "existing.csv"
        with path.open("w", encoding="utf-8", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(headers)
            writer.writerows(rows)
        return path

    def generate(self, **kwargs):
        return self.generator.generate(
            workbook=self.source, output_dir=self.output, mapping_path=self.mapping, **kwargs,
        )

    def test_backend_layout_preserves_hash_utf8_six_steps_and_verbatim_guardrails(self):
        source_before, mapping_before = self.source.read_bytes(), self.mapping.read_bytes()
        with patch.object(self.generator, "hash_uc", side_effect=AssertionError("No salt rotation")):
            self.assertEqual(self.generate(layout="backend"), 1)
        page = PageParser((self.output / "action-synthetic01.html").read_text(encoding="utf-8"))
        self.assertEqual([a["data-etape"] for _, a in page.elements if "data-etape" in a], list("123456"))
        self.assertIn(self.case["guardrails"], page.pre[-1])
        self.assertIn(self.case["description_cas_utilisation"], "".join(page.text))
        self.assertIn(self.case["premiere_action_48h"], "".join(page.text))
        self.assertTrue(any(a.get("href") == "https://avouslia.fr" for _, a in page.elements))
        self.assertTrue(any(a.get("name") == "robots" and "noindex" in a.get("content", "")
                            for _, a in page.elements))
        with (self.output / "mapping_uc_hash.csv").open(encoding="utf-8") as stream:
            self.assertEqual(list(csv.DictReader(stream)), [
                {"case_id": "SYNTHETIC-UX", "case_hash": "synthetic01", "url": "/action-synthetic01.html"},
            ])
        self.assertEqual(self.source.read_bytes(), source_before)
        self.assertEqual(self.mapping.read_bytes(), mapping_before)
        self.assertFalse((self.output / "cases.json").exists())

    def test_directory_layout_and_backend_mapping_input(self):
        self.mapping = self.write_mapping([["SYNTHETIC-UX", "synthetic01"]], ("case_id", "case_hash"))
        self.generate()
        self.assertTrue((self.output / "action" / "synthetic01" / "index.html").is_file())

    def test_leading_audit_and_ambiguous_catalogue_selection(self):
        self.source = self.workbook([("AUDIT", [["notes"]]), ("Sheet1", self.rows)], "leading.xlsx")
        self.assertEqual(len(self.generator.charge_base(self.source)[0]), 1)
        self.source = self.workbook([("A", self.rows), ("B", self.rows)], "ambiguous.xlsx")
        with self.assertRaisesRegex(ValueError, "ambiguous"):
            self.generate()
        self.assertFalse(self.output.exists())
        self.assertEqual(self.generate(sheet_name="B"), 1)

    def test_malformed_sheet1_and_audit_proposal_cannot_be_selected(self):
        for sheets in (
            [("Sheet1", [["bad"]]), ("Other", self.rows)],
            [("BASE_PROPOSEE", self.rows)],
        ):
            self.source = self.workbook(sheets)
            with self.assertRaises(ValueError):
                self.generate()
            self.assertFalse(self.output.exists())
        with self.assertRaisesRegex(ValueError, "Audit proposal"):
            self.generate(sheet_name="BASE_PROPOSEE")

    def test_missing_sheet_headers_duplicate_ids_blank_rag_and_error_cells(self):
        for field, value in (("use_case_id", ""), ("rag_text_auto", ""),
                             ("guardrails", "#VALUE!"), ("guardrails", '="no cache"'),
                             ("mode_execution", "unknown")):
            row = dict(self.record)
            row[field] = value
            self.source = self.workbook([("Sheet1", [self.headers, list(row.values())])])
            with self.subTest(field=field, value=value), self.assertRaises(ValueError):
                self.generate()
            self.assertFalse(self.output.exists())
        for rows in ([self.headers, self.rows[1], self.rows[1]],
                     [self.headers + ["use_case_id"], self.rows[1] + ["SYNTHETIC-UX"]],
                     [["use_case_id"], ["SYNTHETIC-UX"]]):
            self.source = self.workbook([("Sheet1", rows)])
            with self.assertRaises(ValueError):
                self.generate()
        with self.assertRaises(ValueError):
            self.generate(sheet_name="Missing")

    def test_generator_case_insensitive_duplicate_ids_fail_but_original_case_survives(self):
        other = dict(self.record, use_case_id="synthetic-ux")
        self.source = self.workbook([("Sheet1", [*self.rows, list(other.values())])])
        with self.assertRaisesRegex(ValueError, "Duplicate catalogue ID"):
            self.generate()
        self.assertFalse(self.output.exists())
        self.source = self.workbook([("Sheet1", [self.headers, list(other.values())])])
        self.mapping = self.write_mapping([["synthetic-ux", "synthetic01", "/action/synthetic01/"]])
        self.generate(layout="backend")
        with (self.output / "mapping_uc_hash.csv").open(encoding="utf-8") as stream:
            self.assertEqual(next(csv.DictReader(stream))["case_id"], "synthetic-ux")

    def test_development_app_backlink_is_escaped_and_preserves_six_steps(self):
        url = "https://example.invalid/dev-app/?source=parcours&mode=test"
        self.generate(layout="backend", app_url=url)
        html = (self.output / "action-synthetic01.html").read_text(encoding="utf-8")
        page = PageParser(html)
        self.assertTrue(any(a.get("href") == url for _, a in page.elements))
        self.assertIn('href="https://example.invalid/dev-app/?source=parcours&amp;mode=test"', html)
        self.assertEqual([a["data-etape"] for _, a in page.elements if "data-etape" in a], list("123456"))
        self.assertIn(self.case["guardrails"], page.pre[-1])
        self.assertIn("example.invalid/dev-app", "".join(page.text))

    def test_app_backlink_default_environment_and_explicit_precedence(self):
        self.assertEqual(self.generator._validated_app_url(), "https://avouslia.fr")
        with patch.dict(self.generator.os.environ, {"AVOULIA_APP_URL": "https://example.invalid/from-env/"}):
            self.generate(layout="backend")
            self.assertEqual(self.generator._validated_app_url("http://localhost:5173/"), "http://localhost:5173/")
        html = (self.output / "action-synthetic01.html").read_text(encoding="utf-8")
        self.assertIn('href="https://example.invalid/from-env/"', html)

    def test_unsafe_backlinks_fail_before_source_access_or_output(self):
        invalid = (
            "javascript:alert(1)", "data:text/html,test", "file:///test", "ftp://example.invalid",
            "//example.invalid/path", "/relative", "https:///missing-host", "",
            "https://user:pass@example.invalid/", "https://example.invalid/\npath",
            "https://example.invalid\\@other.invalid/", 'https://example.invalid/"onclick="alert(1)',
            "https://example.invalid:99999/", "https://%0aexample.invalid/", "https://[::1",
        )
        for url in invalid:
            with self.subTest(url=url), \
                    patch.object(self.generator, "charge_base", side_effect=AssertionError("No source access")), \
                    self.assertRaisesRegex(ValueError, "HTTP"):
                self.generate(app_url=url)
            self.assertFalse(self.output.exists())
        with self.assertRaises(ValueError):
            self.generator.render_page(self.case, "synthetic01", "synthetic", None, {}, app_url="javascript:alert(1)")

    def test_invalid_mapping_ids_collisions_hashes_and_urls_fail_before_writing(self):
        valid = ["SYNTHETIC-UX", "synthetic01", "/action/synthetic01/"]
        for rows in (
            [valid, valid],
            [valid, ["OTHER", "synthetic01", "/action/synthetic01/"]],
            [["OTHER", "synthetic02", "/action/synthetic02/"]],
            [["", "synthetic01", "/action/synthetic01/"]],
            [["SYNTHETIC-UX", "../escape", "/action/../escape/"]],
            [["SYNTHETIC-UX", "SYNTHETIC01", "/action/SYNTHETIC01/"]],
            [["SYNTHETIC-UX", "synthetic01", "/action/wrong/"]],
            [["SYNTHETIC-UX", "synthetic01", "https://example.invalid/action/synthetic01/"]],
            [valid, ["synthetic-ux", "synthetic02", "/action/synthetic02/"]],
        ):
            self.mapping = self.write_mapping(rows)
            with self.subTest(rows=rows), self.assertRaises(ValueError):
                self.generate()
            self.assertFalse(self.output.exists())

    def test_duplicate_mapping_headers_and_unknown_or_duplicate_selection_fail(self):
        self.mapping = self.write_mapping([], ("use_case_id", "hash", "url", "hash"))
        with self.assertRaisesRegex(ValueError, "headers"):
            self.generate()
        self.mapping = self.write_mapping([["SYNTHETIC-UX", "synthetic01", "/action/synthetic01/"]])
        for targets in (["NOT-IN-WORKBOOK"], ["SYNTHETIC-UX", "SYNTHETIC-UX"]):
            with self.assertRaises(ValueError):
                self.generate(targets=targets)
        self.assertFalse(self.output.exists())

    def test_only_requested_cases_must_have_mapping_and_only_selected_pages_are_written(self):
        other = dict(self.record, use_case_id="UNMAPPED")
        self.source = self.workbook([("Sheet1", [*self.rows, list(other.values())])])
        with self.assertRaisesRegex(ValueError, "missing from mapping"):
            self.generate()
        self.assertEqual(self.generate(targets=["SYNTHETIC-UX"]), 1)
        self.assertEqual(len(list(self.output.rglob("index.html"))), 1)

    def test_historical_mapping_entries_are_retained_but_not_regenerated(self):
        self.mapping = self.write_mapping([
            ["SYNTHETIC-UX", "synthetic01", "/action/synthetic01/"],
            ["HISTORICAL-ONLY", "synthetic02", "/action/synthetic02/"],
        ])
        self.assertEqual(self.generate(layout="backend"), 1)
        self.assertFalse((self.output / "action-synthetic02.html").exists())
        with (self.output / "mapping_uc_hash.csv").open(encoding="utf-8") as stream:
            mapping = list(csv.DictReader(stream))
        self.assertEqual(mapping[1], {
            "case_id": "HISTORICAL-ONLY", "case_hash": "synthetic02", "url": "/action-synthetic02.html",
        })

    def test_replace_dist_refuses_to_remove_historical_mapped_pages(self):
        self.mapping = self.write_mapping([
            ["SYNTHETIC-UX", "synthetic01", "/action/synthetic01/"],
            ["HISTORICAL-ONLY", "synthetic02", "/action/synthetic02/"],
        ])
        self.output.mkdir()
        sentinel = self.output / "action-synthetic02.html"
        sentinel.write_text("historical synthetic page")
        with patch.object(self.generator, "DIST", self.output), self.assertRaisesRegex(ValueError, "historical"):
            self.generate(replace_dist=True)
        self.assertEqual(sentinel.read_text(), "historical synthetic page")

    def test_existing_output_and_missing_output_never_erase(self):
        self.output.mkdir()
        sentinel = self.output / "keep.txt"
        sentinel.write_text("existing")
        with self.assertRaises(ValueError):
            self.generate()
        with self.assertRaises(ValueError):
            self.generate(replace_dist=True)
        with self.assertRaises(ValueError):
            self.generator.generate(workbook=self.source, mapping_path=self.mapping)
        self.assertEqual(sentinel.read_text(), "existing")

    def test_release_needs_mapping_development_is_explicit(self):
        with self.assertRaisesRegex(ValueError, "mapping required"):
            self.generator.generate(workbook=self.source, output_dir=self.output)
        self.assertFalse(self.output.exists())
        with patch.object(self.generator, "hash_uc", return_value="synthetic01"):
            self.assertEqual(self.generator.generate(
                workbook=self.source, output_dir=self.output, development=True,
            ), 1)

    def test_render_and_write_failures_leave_no_partial_output(self):
        with patch.object(self.generator, "render_page", side_effect=ValueError("synthetic render failure")), \
                self.assertRaises(ValueError):
            self.generate()
        self.assertFalse(self.output.exists())
        original = Path.write_text

        def failing_write(path, *args, **kwargs):
            if path.name == "robots.txt":
                raise OSError("synthetic write failure")
            return original(path, *args, **kwargs)

        with patch.object(Path, "write_text", failing_write), self.assertRaises(OSError):
            self.generate()
        self.assertFalse(self.output.exists())
        self.assertEqual(list(self.root.glob(".parcours-*")), [])

    def test_legacy_discovery_rejects_multiple_workbooks(self):
        self.workbook([("Sheet1", self.rows)], "second.xlsx")
        with patch.object(self.generator, "DATA", self.root), self.assertRaises(ValueError):
            self.generator.charge_base()

    def test_explicit_replace_dist_only_after_validation(self):
        self.output.mkdir()
        (self.output / "keep.txt").write_text("existing")
        with patch.object(self.generator, "DIST", self.output):
            with self.assertRaises(ValueError):
                self.generate(replace_dist=True, targets=["UNKNOWN"])
            self.assertTrue((self.output / "keep.txt").is_file())
            self.assertEqual(self.generate(replace_dist=True), 1)
        self.assertFalse((self.output / "keep.txt").exists())
        self.assertTrue((self.output / "action" / "synthetic01" / "index.html").is_file())

    def test_generator_cached_formulas_and_workbook_close(self):
        self.record["rag_text_auto"] = '="synthetic cache"'
        self.source = self.workbook([("Sheet1", [self.headers, list(self.record.values())])])
        from openpyxl.utils import get_column_letter
        coordinate = get_column_letter(self.headers.index("rag_text_auto") + 1) + "2"
        cache_formula(self.source, coordinate, "Texte cache é\nintact")
        books = []

        def opening(*args, **kwargs):
            wb = load_workbook(*args, **kwargs)
            wb.close = Mock(wraps=wb.close)
            books.append(wb)
            return wb

        with patch.object(self.generator, "load_workbook", side_effect=opening):
            frame, _ = self.generator.charge_base(self.source)
        self.assertEqual(frame.iloc[0]["rag_text_auto"], "Texte cache é\nintact")
        for wb in books:
            wb.close.assert_called_once()
        books.clear()
        with patch.object(self.generator, "load_workbook", side_effect=opening), self.assertRaises(ValueError):
            self.generator.charge_base(self.source, "Missing")
        self.assertEqual(len(books), 2)
        for wb in books:
            wb.close.assert_called_once()

    def test_replace_rollback_restores_old_dist_on_failed_promotion(self):
        self.output.mkdir()
        sentinel = self.output / "keep.txt"
        sentinel.write_text("existing")
        original = self.generator._rename_directory

        def fail_promotion(source, destination):
            if source.name.startswith(".parcours-stage-"):
                raise OSError("synthetic promotion failure")
            return original(source, destination)

        with patch.object(self.generator, "DIST", self.output), \
                patch.object(self.generator, "_rename_directory", side_effect=fail_promotion), \
                self.assertRaises(OSError):
            self.generate(replace_dist=True)
        self.assertEqual(sentinel.read_text(), "existing")
        self.assertEqual(list(self.root.glob(".parcours-*")), [])

    def test_cli_explicit_source_succeeds_and_errors_are_non_success_redacted(self):
        arguments = [
            "--workbook", str(self.source), "--sheet", "Sheet1",
            "--mapping", str(self.mapping), "--output-dir", str(self.output),
            "--layout", "backend", "--app-url", "https://example.invalid/cli/", "SYNTHETIC-UX",
        ]
        stdout = io.StringIO()
        with redirect_stdout(stdout):
            self.generator.main(arguments)
        self.assertIn("1 page(s) staged locally", stdout.getvalue())
        self.assertIn("PRIVATE-STAGING", stdout.getvalue())
        self.assertIn("NOT a public webroot", stdout.getvalue())
        self.assertIn('href="https://example.invalid/cli/"',
                      (self.output / "action-synthetic01.html").read_text(encoding="utf-8"))
        stderr = io.StringIO()
        with redirect_stderr(stderr), self.assertRaises(SystemExit) as error:
            self.generator.main(arguments)
        self.assertEqual(error.exception.code, 2)
        self.assertNotIn(str(self.source), stderr.getvalue())
        self.assertNotIn(self.case["guardrails"], stderr.getvalue())


if __name__ == "__main__":
    unittest.main()
